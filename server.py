from __future__ import annotations

import csv
import io
import json
import uuid
from typing import Any

from fastapi import BackgroundTasks, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from backend.app import (
    DB,
    MAX_UPLOAD,
    app,
    process_bytes,
    rows_for_export,
)

# Replace the original upload/export handlers with versions that also support TXT
# input and PDF/TXT output while preserving the existing API contract.
app.router.routes[:] = [
    route
    for route in app.router.routes
    if not (
        getattr(route, "path", None) == "/api/transcricoes"
        and "POST" in getattr(route, "methods", set())
    )
    and not (
        getattr(route, "path", "").startswith("/api/transcricoes/")
        and getattr(route, "path", "").endswith("/planilha")
        and "GET" in getattr(route, "methods", set())
    )
]


def process_text(data: bytes, tipo: str) -> dict[str, Any]:
    text = data.decode("utf-8-sig", errors="replace")
    from backend.app import parse_document
    return parse_document([text], tipo)


def _process_job_extended(transcription_id: str, data: bytes, tipo: str, is_text: bool) -> None:
    try:
        value = process_text(data, tipo) if is_text else process_bytes(data, tipo)
        DB[transcription_id].update(status="concluido", value=value, erro=None)
    except HTTPException as exc:
        DB[transcription_id].update(status="erro", value=None, erro=str(exc.detail))
    except Exception:
        DB[transcription_id].update(status="erro", value=None, erro="Não foi possível processar o documento")


@app.post("/api/transcricoes", status_code=202)
async def create_extended(
    background_tasks: BackgroundTasks,
    arquivo: UploadFile = File(...),
    tipo: str = Form(...),
) -> dict[str, str]:
    if tipo not in {"cartao-ponto", "holerite"}:
        raise HTTPException(400, "Tipo de documento inválido")
    filename = (arquivo.filename or "").lower()
    if not (filename.endswith(".pdf") or filename.endswith(".txt")):
        raise HTTPException(400, "Envie um arquivo PDF ou TXT válido")
    data = await arquivo.read()
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, f"Arquivo excede o limite de {MAX_UPLOAD // 1024 // 1024} MB")
    transcription_id = str(uuid.uuid4())
    is_text = filename.endswith(".txt")
    DB[transcription_id] = {
        "id": transcription_id,
        "tipo": tipo,
        "status": "processando",
        "erro": None,
        "value": None,
    }
    background_tasks.add_task(_process_job_extended, transcription_id, data, tipo, is_text)
    return {"id": transcription_id}


def _download_headers(filename: str) -> dict[str, str]:
    return {"Content-Disposition": f'attachment; filename="{filename}"'}


@app.get("/api/transcricoes/{transcription_id}/planilha")
def export_extended(transcription_id: str, formato: str = "xlsx"):
    item = DB.get(transcription_id)
    if not item:
        raise HTTPException(404, "Transcrição não encontrada")
    if item["status"] != "concluido":
        raise HTTPException(409, "Transcrição ainda não concluída")

    formato = formato.lower()
    headers, rows, warnings = rows_for_export(item["value"], item["tipo"])
    base = f"cartao-de-ponto-{transcription_id}"

    if formato == "json":
        raw = json.dumps(item["value"], ensure_ascii=False, indent=2).encode("utf-8")
        return StreamingResponse(io.BytesIO(raw), media_type="application/json", headers=_download_headers(f"{base}.json"))

    if formato in {"csv", "txt"}:
        out = io.StringIO(newline="")
        writer = csv.writer(out, delimiter="\t" if formato == "txt" else ",", lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)
        if warnings:
            out.write("\nAvisos:\n")
            for idx, messages in enumerate(warnings, start=1):
                if messages:
                    out.write(f"Linha/Página {idx}: {' | '.join(messages)}\n")
        raw = out.getvalue().encode("utf-8-sig")
        media = "text/plain; charset=utf-8" if formato == "txt" else "text/csv; charset=utf-8"
        return StreamingResponse(io.BytesIO(raw), media_type=media, headers=_download_headers(f"{base}.{formato}"))

    if formato == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph

        out = io.BytesIO()
        doc = SimpleDocTemplate(
            out,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
            title="Quick Filler - Cartão de Ponto",
        )
        styles = getSampleStyleSheet()
        story = [Paragraph("Quick Filler — Documento processado", styles["Title"]), Spacer(1, 5 * mm)]
        safe_rows = [[str(x or "") for x in headers]]
        for row in rows:
            safe_rows.append([str(x or "") for x in row])
        if not safe_rows:
            safe_rows = [["Sem dados"]]
        table = Table(safe_rows, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("173772")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("F5F7FA")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(table)
        if any(warnings):
            story.append(Spacer(1, 5 * mm))
            story.append(Paragraph("Avisos de revisão", styles["Heading2"]))
            for idx, messages in enumerate(warnings, start=1):
                if messages:
                    story.append(Paragraph(f"Linha/Página {idx}: {' | '.join(messages)}", styles["BodyText"]))
        doc.build(story)
        out.seek(0)
        return StreamingResponse(out, media_type="application/pdf", headers=_download_headers(f"{base}.pdf"))

    if formato == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side
        wb = Workbook()
        ws = wb.active
        ws.title = "Transcrição"
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="173772")
        warning_fill = PatternFill("solid", fgColor="FFF3CD")
        danger_fill = PatternFill("solid", fgColor="F8D7DA")
        danger_side = Side(style="thin", color="DC3545")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for row_idx, row in enumerate(rows, start=2):
            ws.append(row)
            msgs = warnings[row_idx - 2] if row_idx - 2 < len(warnings) else []
            if msgs:
                fill = danger_fill if any("não sequencial" in m.lower() or "inválida" in m.lower() for m in msgs) else warning_fill
                for cell in ws[row_idx]:
                    cell.fill = fill
                if fill == danger_fill:
                    ws.cell(row_idx, 1).border = Border(left=danger_side)
        for col in ws.columns:
            width = min(max(len(str(c.value or "")) for c in col) + 2, 45)
            ws.column_dimensions[col[0].column_letter].width = width
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=_download_headers(f"{base}.xlsx"))

    raise HTTPException(400, "Formato deve ser xlsx, csv, txt, pdf ou json")


if __name__ == "__main__":
    import os
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
