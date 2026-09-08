from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel
import io, json, uuid, re

app = FastAPI(title='Cartão de Ponto - Quick Filler')
DB = {}
MAX_UPLOAD = 15 * 1024 * 1024

class UpdateBody(BaseModel):
    value: dict

def parse_text(text: str, tipo: str):
    if tipo == 'cartao-ponto':
        days=[]
        for line in text.splitlines():
            m=re.search(r'(\d{2}/\d{2}/\d{4}).*?((?:\d{1,2}:\d{2}\s*)+)', line)
            if not m: continue
            date=m.group(1); times=re.findall(r'\d{1,2}:\d{2}',m.group(2))
            punches=[{'kind':'IN' if i%2==0 else 'OUT','time_raw':t,'time_hhmm':t} for i,t in enumerate(times)]
            days.append({'date_raw':date,'punches':punches})
        return {'pages':[{'page':1,'days':days}]}
    return {'pages':[{'page':1,'year':'','month':'','fields':[],'bases':[]}]}

async def process(upload: UploadFile, tipo: str):
    if tipo not in ('cartao-ponto','holerite'):
        raise HTTPException(400,'Tipo de documento inválido')
    if upload.content_type != 'application/pdf' or not upload.filename.lower().endswith('.pdf'):
        raise HTTPException(400,'Envie um arquivo PDF válido')
    data=await upload.read()
    if len(data)>MAX_UPLOAD: raise HTTPException(413,'PDF excede o limite de 15 MB')
    text=''
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text='\n'.join((p.extract_text() or '') for p in pdf.pages)
    except Exception:
        text=''
    # OCR opcional: se não houver camada de texto, tenta Tesseract quando instalado.
    if not text.strip():
        try:
            import pytesseract
            from pdf2image import convert_from_bytes
            images=convert_from_bytes(data, first_page=1, last_page=20)
            text='\n'.join(pytesseract.image_to_string(img, lang='por+eng') for img in images)
        except Exception:
            pass
    value=parse_text(text,tipo)
    if not value['pages'][0].get('days') and tipo=='cartao-ponto':
        value={'pages':[{'page':1,'days':[]}]}
    return value

@app.get('/', response_class=HTMLResponse)
def home():
    with open('frontend/index.html',encoding='utf-8') as f: return f.read()

@app.get('/healthz')
def health(): return {'status':'ok'}

@app.post('/api/transcricoes', status_code=202)
async def create(arquivo: UploadFile=File(...), tipo: str=Form(...)):
    id=str(uuid.uuid4())
    DB[id]={'id':id,'tipo':tipo,'status':'processando','erro':None,'value':None}
    try: DB[id]['value']=await process(arquivo,tipo); DB[id]['status']='concluido'
    except HTTPException as e: DB[id].update(status='erro',erro=e.detail)
    except Exception: DB[id].update(status='erro',erro='Não foi possível processar o documento')
    return {'id':id}

@app.get('/api/transcricoes/{id}')
def get(id:str):
    if id not in DB: raise HTTPException(404,'Transcrição não encontrada')
    return DB[id]

@app.put('/api/transcricoes/{id}')
def update(id:str, body:UpdateBody):
    if id not in DB: raise HTTPException(404,'Transcrição não encontrada')
    DB[id]['value']=body.value
    return DB[id]

def rows(value,tipo):
    if tipo=='cartao-ponto':
        maxp=max([len(d.get('punches',[])) for p in value.get('pages',[]) for d in p.get('days',[])] or [0])
        out=[]
        for p in value.get('pages',[]):
            for d in p.get('days',[]):
                r={'Data':d.get('date_raw','')}
                for i,x in enumerate(d.get('punches',[]),1): r[f'{"Entrada" if i%2 else "Saída"} {(i+1)//2}']=x.get('time_hhmm','')
                out.append(r)
        return out
    return [{'Pág.':p.get('page',''),'Mês':p.get('month',''),'Ano':p.get('year',''),**{f.get('label',''):f.get('value','') for f in p.get('fields',[])}} for p in value.get('pages',[])]

@app.get('/api/transcricoes/{id}/planilha')
def export(id:str, formato:str='xlsx'):
    if id not in DB: raise HTTPException(404,'Transcrição não encontrada')
    if DB[id]['status']!='concluido': raise HTTPException(409,'Transcrição ainda não concluída')
    formato=formato.lower(); data=rows(DB[id]['value'],DB[id]['tipo'])
    if formato=='json':
        raw=json.dumps(DB[id]['value'],ensure_ascii=False,indent=2).encode(); return StreamingResponse(io.BytesIO(raw),media_type='application/json',headers={'Content-Disposition':f'attachment; filename={id}.json'})
    if formato=='csv':
        import csv
        s=io.StringIO(); fields=list(data[0].keys()) if data else ['Data']; w=csv.DictWriter(s,fieldnames=fields); w.writeheader(); w.writerows(data)
        return StreamingResponse(io.BytesIO(s.getvalue().encode('utf-8-sig')),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename={id}.csv'})
    if formato=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb=Workbook(); ws=wb.active; ws.title='Transcrição'; fields=list(data[0].keys()) if data else ['Data']; ws.append(fields)
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='173772')
        for r in data: ws.append([r.get(f,'') for f in fields])
        out=io.BytesIO(); wb.save(out); out.seek(0); return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename={id}.xlsx'})
    raise HTTPException(400,'Formato deve ser xlsx, csv ou json')
